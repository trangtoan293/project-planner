"""check_plan.py phải ĐẠT với plan sạch và bắt đúng các lỗi lẫn dữ liệu cũ.
Chạy: python3 tests/test_check_plan.py
"""
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HEADER = "wbs,name,duration,start,predecessors,pm_n,pm_pct,ba_n,ba_pct,dev_n,dev_pct,syse_n,syse_pct,tester_n,tester_pct\n"
PARAMS = """customer: KHA
kickoff: 2026-11-02
deadline: 2026-12-31
holidays: 2026-11-20
team_max: PM=1,BA=1,DEV=2,SysE=1,Tester=1
proposal_md: 48
milestones: Go-live<=2026-12-18
stale_terms: KHB, KHC, KHD
"""
# Task ID = thứ tự dòng: 1.1 là task 2, 1.2 (sau task 2) là task 3, 1.3 (sau task 3); 20/11 là ngày lễ → 1.2 kết thúc 23/11
CLEAN = HEADER + """1,Dự án KHA,,,,,,,,,,,,,
1.1,Khởi động,5,2026-11-02,,1,1,1,1,0,0,0,0,0,0
1.2,Xây dựng,10,2026-11-09,2,1,0.2,0,0,2,1,1,0.5,1,0.5
1.3,Go-live và nghiệm thu,1,2026-11-24,3,1,1,1,1,2,1,1,1,1,1
"""


def run(params, csv, proposal=""):
    with tempfile.TemporaryDirectory() as d:
        p, c, m = Path(d) / "params.txt", Path(d) / "wbs.csv", Path(d) / "proposal.md"
        p.write_text(params, encoding="utf-8")
        c.write_text(csv, encoding="utf-8")
        m.write_text(proposal, encoding="utf-8")
        r = subprocess.run([sys.executable, ROOT / "scripts" / "check_plan.py", p, c, m], capture_output=True, text=True)
        return r.returncode, r.stdout + r.stderr


def expect_fail(csv, needle, params=PARAMS, proposal=""):
    code, out = run(params, csv, proposal)
    assert code == 1 and needle in out, f"cần FAIL chứa '{needle}', được:\n{out}"


def main():
    code, out = run(PARAMS, CLEAN, "Khởi động 02/11/2026, go-live 24/11/2026.")
    assert code == 0, out

    expect_fail(CLEAN.replace("2026-11-02", "2025-11-03"), "trước kickoff")             # ngày của dự án cũ
    expect_fail(CLEAN.replace("2026-11-24", "2027-01-05"), "sau deadline")
    # hypercare/bảo hành được phép sau deadline khi khai báo post_deadline; thiếu khai báo thì vẫn FAIL
    support = CLEAN + "1.4,Hypercare 1 tháng,10,2027-01-04,4,0,0,0,0,0,0,0,0,0,0\n"
    expect_fail(support, "sau deadline")
    code, out = run(PARAMS + "post_deadline: Hypercare; Bảo hành\n", support, "Khởi động 02/11/2026, go-live 24/11/2026.")
    assert code == 0, out
    expect_fail(CLEAN.replace("2026-11-09,2", "2026-11-04,2"), "sớm hơn ràng buộc 2")    # đè lên predecessor
    expect_fail(CLEAN.replace("2026-11-09", "2026-11-08"), "ngày nghỉ")                  # Chủ nhật
    expect_fail(CLEAN.replace("2026-11-24,3", "2026-11-20,3"), "ngày nghỉ")              # ngày lễ trong params
    expect_fail(CLEAN.replace(",2,1,1,0.5,1,0.5", ",3,1,1,0.5,1,0.5"), "team_max")       # team size cũ
    expect_fail(CLEAN, "lệch > 5%", params=PARAMS.replace("proposal_md: 48", "proposal_md: 845"))
    expect_fail(CLEAN, "Go-live", params=PARAMS.replace("Go-live<=2026-12-18", "Go-live<=2026-11-20"))
    expect_fail(CLEAN.replace("Dự án KHA", "Dự án KHB"), "'KHB'")                        # tên khách cũ
    expect_fail(CLEAN, "'KHC'", proposal="Như dự án KHC, chúng tôi...")

    sys.path.insert(0, str(ROOT / "scripts"))
    from check_plan import check_requirements
    prop = "# 5. THU THẬP\n## 5.2. IGT\n# PHỤ LỤC B. TIÊU CHÍ\n"
    req = "| R01 | a | Câu hỏi | §7 | Mục 5.2 | x |\n| R02 | b | Kết quả | §8 | Phụ lục B | y |\n"
    assert check_requirements(req, prop) == []
    assert "khoảng" in check_requirements("| R03 | c | Kết quả | §8 | Mục 5–11 | z |", prop)[0]
    assert "không có mục '6.1'" in check_requirements("| R04 | d | Kết quả | §8 | Mục 6.1 | z |", prop)[0]
    assert "chưa có" in check_requirements("| R05 | e | Kết quả | §8 | | |", prop)[0]
    prop_b = "# II. CÁCH TIẾP CẬN\n## II.1. Giải pháp\n### II.1.2. Kiến trúc\n# PHỤ LỤC A. CÂU HỎI\n"  # loại B, số La Mã
    assert check_requirements("| R06 | f | Chức năng | 2.1 | Mục II.1.2 | z |\n| R07 | g | Kết quả | 2.2 | Phụ lục A | z |", prop_b) == []
    assert "không có mục 'II.3'" in check_requirements("| R08 | h | Kết quả | 2.2 | Mục II.3 | z |", prop_b)[0]

    from check_plan import check_summary
    template = (ROOT / "templates" / "01-requirements.md").read_text(encoding="utf-8")
    assert len(check_summary(template)) == 9                                  # template rỗng: 9 mục trống
    filled = "\n".join(f"### {i}. {name}\n- ý (R01) [T]" for i, name in enumerate(
        ["Mục tiêu dự án", "Yêu cầu dự án", "Phạm vi", "Công nghệ yêu cầu", "Kết quả mong muốn",
         "Nên làm", "Không nên làm", "Tiêu chí đánh giá", "Khoảng trống và điểm cần làm rõ"], 1))
    assert check_summary(filled) == []
    real = filled.replace("- ý (R01) [T]", "#### 2.1 Con\n- ý một\n- ý hai") + "\n## Thuật ngữ\n| a | b |\n"
    assert check_summary(real) == []                                          # mục nhiều dòng, có phần sau
    assert "mục 6. Nên làm còn trống" in check_summary(real.replace("### 6. Nên làm\n#### 2.1 Con\n- ý một\n- ý hai", "### 6. Nên làm\n"))[0]
    assert "thiếu mục 'Nên làm'" in check_summary(filled.replace("### 6. Nên làm", "### 6. Khác"))[0]

    test_refs_and_team()
    test_sprints()
    print("OK")


def test_sprints():
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        p, c, out = Path(d) / "params.txt", Path(d) / "wbs.csv", Path(d) / "plan.xlsx"
        p.write_text(PARAMS, encoding="utf-8")
        c.write_text(CLEAN, encoding="utf-8")
        subprocess.run([sys.executable, ROOT / "scripts" / "build_master_plan.py", c, out, "--params", p],
                       capture_output=True, text=True, check=True)
        wb = openpyxl.load_workbook(out)
        plan, sp = wb["Master_Plan"], wb["Sprint"]
        assert plan["AA2"].value == "Sprint" and plan["AA4"].value == "S1" and plan["AA5"].value == "S1–S2", \
            [plan[f"AA{r}"].value for r in range(2, 7)]                                 # 1.2: 09/11 → 23/11
        rows = [r for r in sp.iter_rows(min_row=3, values_only=True) if r[0]]
        assert [r[0] for r in rows] == ["S1", "S2"], rows                               # kickoff 02/11 → 24/11
        assert rows[0][3] == 10 and rows[1][3] == 9, rows                               # 20/11 là ngày lễ
        assert abs(sum(r[9] for r in rows) - 48) < 0.01, rows                           # tổng MD các sprint = plan
        assert rows[0][13] == "OK"
        c.write_text(CLEAN.replace("1.2,Xây dựng,10,2026-11-09,2,1,0.2,0,0,2,1", "1.2,Xây dựng,10,2026-11-09,2,1,0.2,0,0,5,1"),
                     encoding="utf-8")                                                  # S1: DEV 5 người × 5 ngày = 25 MD > 2 × 10
        subprocess.run([sys.executable, ROOT / "scripts" / "build_master_plan.py", c, out, "--params", p],
                       capture_output=True, text=True, check=True)
        rows = [r for r in openpyxl.load_workbook(out)["Sprint"].iter_rows(min_row=3, values_only=True) if r[0]]
        assert rows[0][13] == "VƯỢT TẢI", rows


REQ = """| **Mã** | **Trích yêu cầu** | **Loại** | **Nguồn** | **Trả lời tại** | **Trả lời tóm tắt** |
| R01 | Nạp dữ liệu T24 | Chức năng | §2 | Mục 5.1 | x |
| R02 | Đào tạo 20 học viên | Dịch vụ | §3 | Mục 12 | x |
| R03 | Tên dự án | Bối cảnh | §1 | Mục 1.1 | x |
| R04 | Tích hợp JIRA | Chức năng | §4 | Mục 2.3 | Ngoài phạm vi, change request |
"""
TEAM_CSV = """wbs,name,predecessors,owner,refs,note
1,Dự án KHA,,,,
1.1,M0 — Nền tảng,,,,Dev/UAT chạy được
1.1.1,Cài Dev/UAT,,DevOps,,
1.2,M1 — Lát T24,,,R01,Báo cáo 1 chạy trên UAT; biên bản demo
1.2.1,Nạp T24 vào Raw,1.1.1,DE,R01,
1.2.2,Đào tạo,1.2.1,BA,R02,
"""


def test_refs_and_team():
    from check_plan import check_refs
    tasks = [{"wbs": "1", "refs": []}, {"wbs": "1.1.1", "refs": ["R01"]}]
    fails, _ = check_refs(REQ, tasks)
    assert fails == ["R02 (Dịch vụ) chưa thuộc milestone nào (thêm vào cột refs của việc làm nó)"], fails  # R03 bối cảnh, R04 ngoài phạm vi: bỏ qua
    assert check_refs(REQ, tasks + [{"wbs": "1.2", "refs": ["R02"]}]) == ([], [])
    assert "refs 'R99' không có" in check_refs(REQ, tasks + [{"wbs": "1.3", "refs": ["R02", "R99"]}])[0][0]
    assert "chưa có cột refs" in check_refs(REQ, [{"wbs": "1", "refs": None}])[1][0]          # CSV cũ: chỉ WARN

    with tempfile.TemporaryDirectory() as d:
        c, out = Path(d) / "03-wbs.csv", Path(d) / "03-task-list.xlsx"
        c.write_text(TEAM_CSV, encoding="utf-8")
        (Path(d) / "01-requirements.md").write_text(REQ, encoding="utf-8")
        cmd = [sys.executable, ROOT / "scripts" / "build_master_plan.py", c, out, "--team"]
        r = subprocess.run(cmd, capture_output=True, text=True)
        assert r.returncode == 0 and "REFS: ĐẠT" in r.stdout, r.stdout + r.stderr
        import openpyxl
        ws = openpyxl.load_workbook(out).active
        head = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
        assert not any(h and ("MD" in h or "Start" in h or "Duration" in h) for h in head), head  # không ngày, không MD
        rows = {ws.cell(i, 1).value: [ws.cell(i, j).value for j in range(1, 8)] for i in range(2, ws.max_row + 1)}
        assert rows["1.2.1"][3:6] == ["DE", "3", "R01"], rows["1.2.1"]                       # vai trò, phụ thuộc theo Task ID, mã R
        assert rows["1.2"][6] == "Báo cáo 1 chạy trên UAT; biên bản demo"
        c.write_text(TEAM_CSV.replace("1.2.1,BA,R02", "1.2.1,BA,"), encoding="utf-8")       # bỏ R02 → thiếu
        r = subprocess.run(cmd, capture_output=True, text=True)
        assert r.returncode == 1 and "R02 (Dịch vụ) chưa thuộc milestone" in r.stdout, r.stdout


if __name__ == "__main__":
    main()
