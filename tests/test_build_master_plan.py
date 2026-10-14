"""Round-trip: master plan mẫu (198 MD) → CSV → build_master_plan.py → so khớp MD và công thức.
Chạy: python3 tests/test_build_master_plan.py
"""
import csv
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "samples/master_plan_sample" / "master_plan_Phase1_Search_ChatDocs.xlsx"
COLS = [("pm", "K", "L"), ("ba", "N", "O"), ("dev", "Q", "R"), ("syse", "T", "U"), ("tester", "W", "X")]


def sample_to_csv(path):
    ws = openpyxl.load_workbook(SAMPLE).worksheets[0]
    fields = ["wbs", "name", "duration", "start", "predecessors"] + [f"{k}_{s}" for k, *_ in COLS for s in ("n", "pct")]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fields)
        w.writeheader()
        for r in range(3, 64):
            if not ws[f"A{r}"].value:
                continue
            row = {"wbs": ws[f"A{r}"].value, "name": ws[f"C{r}"].value}
            if ws[f"K{r}"].value is not None:
                row["duration"] = ws[f"D{r}"].value.split()[0]
                row["start"] = datetime.strptime(ws[f"E{r}"].value[4:], "%m/%d/%y").date().isoformat()
                row["predecessors"] = ws[f"G{r}"].value or ""
                for k, cn, cp in COLS:
                    row[f"{k}_n"], row[f"{k}_pct"] = ws[f"{cn}{r}"].value, ws[f"{cp}{r}"].value
            w.writerow(row)


def main():
    with tempfile.TemporaryDirectory() as d:
        src, out = Path(d) / "wbs.csv", Path(d) / "plan.xlsx"
        sample_to_csv(src)
        res = subprocess.run([sys.executable, ROOT / "scripts" / "build_master_plan.py", src, out],
                             capture_output=True, text=True, check=True)
        assert "~198.0 MD" in res.stdout, res.stdout

        ws = openpyxl.load_workbook(out).worksheets[0]
        assert ws["A2"].value == "WBS" and ws["Z2"].value == "Kiểm tra"
        assert ws["M3"].value == "=SUM(M4,M14,M22,M37,M47,M60)", ws["M3"].value  # tổng dự án = các milestone
        assert ws["M5"].value == "=SUM(M6,M7)", ws["M5"].value                  # nhóm cấp 3 = các việc con
        assert ws["M6"].value == "=$J6*K6*L6"
        assert ws["D6"].value == "1 day" and ws["E6"].value == "Tue 10/13/26"
        assert ws["Q1"].value == 3                                               # DEV tối đa 3 người
        assert ws["B66"].value == "TỔNG HỢP NGUỒN LỰC VÀ KIỂM TRA"
        assert ws["C77"].value == "=IF(C76>=C75,\"ĐẠT\",\"KHÔNG ĐẠT\")", ws["C77"].value  # giống mẫu
    print("OK")


if __name__ == "__main__":
    if not SAMPLE.exists():  # mẫu thật không đưa lên git
        sys.exit(f"bỏ qua: thiếu {SAMPLE.relative_to(ROOT)}")
    main()
