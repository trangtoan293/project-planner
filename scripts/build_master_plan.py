"""Sinh master plan .xlsx đúng format mẫu (samples/master_plan_sample/) từ một file WBS dạng CSV.

Dùng:
    python3 scripts/build_master_plan.py projects/<kh>/03-wbs.csv projects/<kh>/04-master-plan.xlsx \
        --params projects/<kh>/00-params.txt
    (hoặc truyền tay: --holidays 2027-02-04:2027-02-10,2027-04-30 --pm-min 0.8 --previous-total 1640.04)

CSV (UTF-8, có header):
    wbs,name,duration,start,predecessors,pm_n,pm_pct,ba_n,ba_pct,dev_n,dev_pct,syse_n,syse_pct,tester_n,tester_pct
- Dòng tổng hợp (có con theo tiền tố WBS) chỉ cần wbs + name; ngày và MD tự cộng từ con.
- Dòng chi tiết: duration = số ngày làm việc, *_pct là 0..1.
- start: ĐỂ TRỐNG để script tự tính từ predecessors (hoặc kickoff nếu không có predecessor).
  Chỉ ghi YYYY-MM-DD khi ngày đó bị khách/hợp đồng cố định.
- predecessors: mã WBS hoặc Task ID, cách nhau dấu phẩy; hỗ trợ FS (mặc định), SS, FF và lag: 1.2.1, 1.3.2SS+2, 1.4.1FF.
"""
import argparse
import csv
import re
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# (tên hiển thị, khóa CSV, cột số người, cột %, cột MD, nhãn MD) — giữ đúng thứ tự cột của file mẫu
ROLES = [
    ("PM", "pm", "K", "L", "M", "PM"),
    ("BA", "ba", "N", "O", "P", "BA"),
    ("DEV", "dev", "Q", "R", "S", "SE"),
    ("SysE", "syse", "T", "U", "V", "SysE"),
    ("Tester", "tester", "W", "X", "Y", "Tester"),
]
ROLE_FULL = ["Project Manager", "Business Analyst", "Developer", "System Engineer", "Tester"]
HEADERS = ["WBS", "Task ID", "Task Name", "Duration", "Start", "Finish", "Predecessors",
           "Resource Names", "Tổng MD", "Số ngày\n(theo Duration)"]
NOTE = ("Ô chữ xanh = ô nhập liệu (Số người & % phân bổ tại dòng chi tiết; số người tối đa dòng 1; "
        "tham số cuối bảng). MD và mọi cột còn lại là công thức: MD = Số ngày × Số người × % phân bổ.")

FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = "FF0000FF"
LEVEL_STYLE = {1: ("FFD9E2F3", "FF000000"), 2: ("FFEEF3FA", "FF000000"), 3: ("FFD9E2F3", "FF1F3864")}
WIDTHS = {"A": 8, "B": 9.7, "C": 46, "D": 9, "E": 14.4, "F": 14, "G": 12, "H": 17, "I": 8, "J": 9, "Z": 12}


def parse_holidays(spec):
    days = set()
    for part in filter(None, (p.strip() for p in (spec or "").split(","))):
        a, _, b = (s.strip() for s in part.partition(":"))
        d, end = date.fromisoformat(a), date.fromisoformat(b or a)
        while d <= end:
            days.add(d)
            d += timedelta(days=1)
    return days


def is_workday(d, holidays):
    return d.weekday() < 5 and d not in holidays


def add_workdays(start, n, holidays):
    """Ngày kết thúc của công việc dài n ngày làm việc bắt đầu từ start."""
    d = start
    while not is_workday(d, holidays):
        d += timedelta(days=1)
    for _ in range(n - 1):
        d += timedelta(days=1)
        while not is_workday(d, holidays):
            d += timedelta(days=1)
    return d


def count_workdays(start, finish, holidays):
    return sum(is_workday(start + timedelta(days=i), holidays) for i in range((finish - start).days + 1))


def fmt_date(d):
    return f"{d.strftime('%a')} {d.month}/{d.day}/{d.strftime('%y')}"  # kiểu MS Project: Fri 9/4/26


def fmt_duration(n):
    return f"{n} day" if n == 1 else f"{n} days"


def num(v, field, wbs):
    try:
        return float(v) if v not in (None, "") else 0.0
    except ValueError:
        sys.exit(f"WBS {wbs}: giá trị '{v}' ở cột {field} không phải số")


PRED_RE = re.compile(r"^([\d.]+?)\s*(FS|SS|FF)?\s*([+-]\s*\d+)?\s*(?:days?|d)?$", re.I)


def shift_workdays(d, n, holidays):
    """Lùi/tiến n ngày làm việc (n=0 giữ nguyên)."""
    step = 1 if n >= 0 else -1
    for _ in range(abs(n)):
        d += timedelta(days=step)
        while not is_workday(d, holidays):
            d += timedelta(days=step)
    return d


def parse_preds(t, tasks):
    """'1.2.1', '12', '1.2.1SS+2', '48FF' → [(task, kind, lag, raw)]. Số có dấu chấm là mã WBS, số nguyên là Task ID."""
    by_wbs = {x["wbs"]: x for x in tasks}
    out = []
    for raw in filter(None, (p.strip() for p in t["pred"].split(","))):
        m = PRED_RE.match(raw)
        if not m:
            sys.exit(f"WBS {t['wbs']}: không đọc được predecessor '{raw}' (dạng hợp lệ: 1.2.1, 12, 1.2.1SS+2, 48FF)")
        ref = m.group(1)
        p = by_wbs.get(ref) if "." in ref else (tasks[int(ref) - 1] if 0 < int(ref) <= len(tasks) else None)
        if p is None:
            sys.exit(f"WBS {t['wbs']}: predecessor '{ref}' không tồn tại")
        if p["parent"]:
            sys.exit(f"WBS {t['wbs']}: predecessor '{ref}' là dòng tổng hợp, hãy trỏ tới việc chi tiết")
        out.append((p, (m.group(2) or "FS").upper(), int((m.group(3) or "0").replace(" ", "")), raw))
    return out


def earliest_start(t, preds, holidays):
    """Ngày bắt đầu sớm nhất thỏa mọi ràng buộc predecessor."""
    starts = []
    for p, kind, lag, _ in preds:
        if kind == "FS":
            starts.append(shift_workdays(p["finish"], 1 + lag, holidays))
        elif kind == "SS":
            starts.append(shift_workdays(p["start"], lag, holidays))
        else:  # FF: kết thúc không sớm hơn predecessor + lag
            starts.append(shift_workdays(shift_workdays(p["finish"], lag, holidays), -(t["days"] - 1), holidays))
    return max(starts)


def schedule(tasks, holidays, kickoff):
    """Điền start/finish cho việc chi tiết: start ghi tay > theo predecessor > kickoff."""
    leaves = [t for t in tasks if not t["parent"]]
    pending = list(leaves)
    while pending:
        progressed = False
        for t in list(pending):
            preds = parse_preds(t, tasks)
            if any("finish" not in p for p, *_ in preds):
                continue
            if t["start_fixed"]:
                t["start"] = t["start_fixed"]  # giữ nguyên kể cả khi sai, để check_plan.py báo lỗi
            elif preds:
                t["start"] = earliest_start(t, preds, holidays)
            elif kickoff:
                t["start"] = add_workdays(kickoff, 1, holidays)
            else:
                sys.exit(f"WBS {t['wbs']}: không có start, predecessor hay kickoff (--params) để tính ngày")
            t["finish"] = add_workdays(t["start"], t["days"], holidays)
            pending.remove(t)
            progressed = True
        if not progressed:
            sys.exit("Predecessor bị vòng lặp: " + ", ".join(t["wbs"] for t in pending))


def load_tasks(path, holidays, kickoff=None):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.DictReader(f) if (r.get("wbs") or "").strip()]
    if not rows:
        sys.exit("CSV không có dòng nào")
    wbs_list = [r["wbs"].strip() for r in rows]
    if len(set(wbs_list)) != len(wbs_list):
        sys.exit("CSV có WBS trùng")
    tasks = []
    for r in rows:
        wbs = r["wbs"].strip()
        t = {"wbs": wbs, "name": r["name"].strip(), "pred": (r.get("predecessors") or "").strip(),
             "level": wbs.count(".") + 1,
             "parent": any(w.startswith(wbs + ".") for w in wbs_list)}
        if not t["parent"]:
            t["days"] = int(num(r.get("duration"), "duration", wbs))
            if t["days"] < 1:
                sys.exit(f"WBS {wbs}: dòng chi tiết cần duration >= 1")
            start = (r.get("start") or "").strip()
            t["start_fixed"] = date.fromisoformat(start) if start else None
            for _, key, *_ in ROLES:
                n, pct = num(r.get(f"{key}_n"), f"{key}_n", wbs), num(r.get(f"{key}_pct"), f"{key}_pct", wbs)
                if n < 0 or not 0 <= pct <= 1:
                    sys.exit(f"WBS {wbs}: {key}_n phải >= 0 và {key}_pct trong khoảng 0..1")
                t[key] = (n, pct)
        tasks.append(t)
    schedule(tasks, holidays, kickoff)
    # dòng tổng hợp: ngày = min/max của các dòng chi tiết bên dưới
    for t in tasks:
        if t["parent"]:
            leaves = [c for c in tasks if not c["parent"] and c["wbs"].startswith(t["wbs"] + ".")]
            t["start"] = min(c["start"] for c in leaves)
            t["finish"] = max(c["finish"] for c in leaves)
            t["days"] = count_workdays(t["start"], t["finish"], holidays)
    return tasks


def style(cell, fill=None, bold=False, color="FF000000", size=10, fmt=None, align="center", wrap=None):
    cell.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt:
        cell.number_format = fmt


def write_header(ws, max_people):
    ws["A1"] = NOTE
    ws["A1"].font = Font(name=FONT, size=9, color=BLUE)
    ws["J1"] = "Số người tối đa →"
    ws["J1"].font = Font(name=FONT, size=9, bold=True)
    ws["J1"].alignment = Alignment(horizontal="right")
    for (_, key, cn, *_ ) in ROLES:
        ws[f"{cn}1"] = max_people[key]
        style(ws[f"{cn}1"], color=BLUE, fmt="0")
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(2, i, h)
        style(c, fill="FF1F4E79", bold=True, color="FFFFFFFF", wrap=True)
    for disp, _, cn, cp, cm, md_label in ROLES:
        for col, text in ((cn, f"{disp}\nSố\nngười"), (cp, f"{disp}\n%\nphân bổ"), (cm, f"{md_label}\nMD")):
            style(ws[f"{col}2"], fill="FF2E75B6", bold=True, color="FFFFFFFF", size=9, wrap=True)
            ws[f"{col}2"] = text
    ws["Z2"] = "Kiểm tra"
    style(ws["Z2"], fill="FF1F4E79", bold=True, color="FFFFFFFF", wrap=True)
    ws.row_dimensions[2].height = 43.5


def write_task(ws, r, idx, t, child_rows):
    resources = ",".join(f'IF({cm}{r}>0,"{label}","")' for _, _, _, _, cm, label in ROLES)
    values = {
        "A": t["wbs"], "B": idx, "C": t["name"], "D": fmt_duration(t["days"]),
        "E": fmt_date(t["start"]), "F": fmt_date(t["finish"]), "G": t.get("pred_ids") or None,
        "H": f'=_xlfn.TEXTJOIN(", ",TRUE(),{resources})',
        "I": "=" + "+".join(f"{cm}{r}" for *_, cm, _ in ROLES),
        "J": f'=IF($D{r}="","",VALUE(LEFT($D{r},FIND(" ",$D{r})-1)))',
    }
    for _, key, cn, cp, cm, _ in ROLES:
        if t["parent"]:
            values[cm] = "=SUM(" + ",".join(f"{cm}{c}" for c in child_rows) + ")"
        else:
            values[cn], values[cp] = t[key]
            values[cm] = f"=$J{r}*{cn}{r}*{cp}{r}"
    if not t["parent"]:
        over = ",".join(f"{cn}{r}>{cn}$1,{cp}{r}>1" for _, _, cn, cp, *_ in ROLES)
        values["Z"] = f'=IF(I{r}=0,"Chưa phân bổ",IF(OR({over}),"VƯỢT TẢI","OK"))'

    fill, color = LEVEL_STYLE.get(t["level"], (None, "FF000000")) if t["parent"] else (None, "FF000000")
    inputs = {c for _, _, cn, cp, *_ in ROLES for c in (cn, cp)}
    pct_cols = {cp for _, _, _, cp, *_ in ROLES}
    for col in [get_column_letter(i) for i in range(1, 27)]:
        cell = ws[f"{col}{r}"]
        if col in values:
            cell.value = values[col]
        fmt = "0.0" if col in {"I"} | {cm for *_, cm, _ in ROLES} else "0%" if col in pct_cols else "0" if col in inputs else None
        style(cell, fill=fill, bold=t["parent"], color=BLUE if (col in inputs and not t["parent"]) else color,
              fmt=fmt, align="left" if col == "C" else "center", wrap=True if col == "C" else None)


def write_summary(ws, r0, pm_min, previous_total):
    hdr = dict(fill="FF1F4E79", bold=True, color="FFFFFFFF", wrap=True)
    ws[f"B{r0}"] = "TỔNG HỢP NGUỒN LỰC VÀ KIỂM TRA"
    style(ws[f"B{r0}"], **hdr)
    for i, h in enumerate(["Vai trò", "Số người tối đa", "Tổng MD", "Tỷ trọng trên tổng",
                           "Năng lực tối đa\n(người × ngày dự án)", "Kiểm tra năng lực"]):
        c = ws.cell(r0 + 1, 2 + i, h)
        style(c, **hdr)
    for i, ((_, _, cn, _, cm, _), full) in enumerate(zip(ROLES, ROLE_FULL)):
        r = r0 + 2 + i
        row = [full, f"={cn}1", f"={cm}3", f"={cm}3/$I$3", f"={cn}1*$J$3", f'=IF(D{r}<=F{r},"ĐẠT","VƯỢT")']
        for j, (v, fmt) in enumerate(zip(row, [None, "0", "0.0", "0.0%", "0", None])):
            c = ws.cell(r, 2 + j, v)
            style(c, fmt=fmt, align="left" if j == 0 else "center")
    rt = r0 + 2 + len(ROLES)
    for col, v, fmt in (("B", "TỔNG CỘNG", None), ("D", "=I3", "0.0"), ("E", f"=SUM(E{r0+2}:E{rt-1})", "0.0%")):
        ws[f"{col}{rt}"] = v
        style(ws[f"{col}{rt}"], bold=True, fmt=fmt, align="left" if col == "B" else "center")

    params = []
    if previous_total is not None:
        params += [("Tổng MD phương án đã gửi khách hàng", previous_total, "0.00", True),
                   ("Tổng MD phương án mới", "=I3", "0.00", False),
                   ("Kiểm tra: Tổng mới không thấp hơn tổng cũ", "=IF(C{p1}>=C{p0},\"ĐẠT\",\"KHÔNG ĐẠT\")", None, False)]
    params += [("Mức tham gia tối thiểu của PM trên thời gian dự án (tham số, điều chỉnh được)", pm_min, "0%", True),
               ("Mức tham gia thực tế của PM = MD PM ÷ (Số ngày dự án × Số PM)", "=M3/($J$3*$K$1)", "0%", False),
               ("Kiểm tra: PM đạt mức tham gia tối thiểu", "=IF(C{p1}>=C{p0},\"ĐẠT\",\"KHÔNG ĐẠT\")", None, False)]
    r = rt + 2
    for i, (label, v, fmt, is_input) in enumerate(params):
        if isinstance(v, str) and "{p0}" in v:
            v = v.format(p0=r - 2, p1=r - 1)
        ws[f"B{r}"], ws[f"C{r}"] = label, v
        style(ws[f"B{r}"], align="left", wrap=True)
        style(ws[f"C{r}"], color=BLUE if is_input else "FF000000", fmt=fmt)
        r += 1


def build(tasks, out, sheet_title, pm_min, previous_total):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    leaves = [t for t in tasks if not t["parent"]]
    max_people = {key: int(max(t[key][0] for t in leaves)) for _, key, *_ in ROLES}
    write_header(ws, max_people)

    row_of = {t["wbs"]: 3 + i for i, t in enumerate(tasks)}
    task_id = {t["wbs"]: i + 1 for i, t in enumerate(tasks)}
    for t in leaves:  # cột Predecessors theo Task ID như file MS Project
        t["pred_ids"] = ",".join(
            f"{task_id[p['wbs']]}{'' if kind == 'FS' and not lag else kind}"
            f"{'' if not lag else f'{lag:+d} day' + ('s' if abs(lag) > 1 else '')}"
            for p, kind, lag, _ in parse_preds(t, tasks))
    for i, t in enumerate(tasks):
        children = [row_of[c["wbs"]] for c in tasks
                    if c["wbs"].startswith(t["wbs"] + ".") and c["level"] == t["level"] + 1]
        write_task(ws, row_of[t["wbs"]], i + 1, t, children)

    write_summary(ws, 3 + len(tasks) + 2, pm_min, previous_total)
    for col in [get_column_letter(i) for i in range(1, 27)]:
        ws.column_dimensions[col].width = WIDTHS.get(col, 8 if col in {"L", "O", "R", "U", "X"} else 6 if col in {"K", "N", "Q", "T", "W"} else 7)
    ws.freeze_panes = "D3"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv")
    ap.add_argument("out")
    ap.add_argument("--sheet", default="Master_Plan")
    ap.add_argument("--holidays", default="", help="YYYY-MM-DD hoặc khoảng YYYY-MM-DD:YYYY-MM-DD, cách nhau dấu phẩy")
    ap.add_argument("--pm-min", type=float, default=0.8)
    ap.add_argument("--previous-total", type=float, default=None, help="Tổng MD bản đã gửi khách (để kiểm tra không giảm)")
    ap.add_argument("--params", help="00-params.txt của dự án: lấy holidays, pm_min, previous_total từ đây")
    a = ap.parse_args()
    if a.params:
        from check_plan import read_params  # cùng thư mục scripts/
        p = read_params(a.params)
        a.holidays = a.holidays or p.get("holidays", "")
        a.pm_min = float(p.get("pm_min", a.pm_min))
        a.previous_total = a.previous_total or (float(p["previous_total"]) if p.get("previous_total") else None)
        kickoff = date.fromisoformat(p["kickoff"])
    else:
        kickoff = None
    tasks = load_tasks(a.csv, parse_holidays(a.holidays), kickoff)
    if tasks[0]["level"] != 1 or sum(t["level"] == 1 for t in tasks) != 1:
        sys.exit("Dòng đầu tiên phải là dòng tổng dự án (WBS '1') và là dòng cấp 1 duy nhất")
    build(tasks, a.out, a.sheet, a.pm_min, a.previous_total)
    total = sum(t["days"] * n * p for t in tasks if not t["parent"] for _, key, *_ in ROLES for n, p in [t[key]])
    print(f"Đã ghi {a.out}: {len(tasks)} dòng, {tasks[0]['days']} ngày làm việc, ~{total:.1f} MD")


if __name__ == "__main__":
    main()
