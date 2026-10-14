"""Chuyển proposal markdown → .docx theo format proposal mẫu (Roboto Condensed, bảng tiêu đề xanh, bìa, mục lục).

Dùng:
    python3 scripts/md2docx.py projects/<kh>/02-proposal.md projects/<kh>/02-proposal.docx
    python3 scripts/md2docx.py projects/<kh>/02-proposal.md projects/<kh>/02-proposal.docx --draft   # bản nháp nội bộ

Bản gửi khách chỉ xuất được khi 02-review.md (cùng thư mục) có dòng "KẾT LUẬN: ĐẠT".
--draft bỏ qua kiểm tra và ghi ra <tên>-NHAP.docx, không dùng để gửi khách.

Front matter ở đầu file markdown (giữa hai dòng ---) điền trang bìa và header:
    vendor: KATALYST
    title: ĐỀ XUẤT GIẢI PHÁP VÀ PHƯƠNG ÁN TRIỂN KHAI
    subtitle: KHA DATA WAREHOUSE TRÊN HANAS DATA PLATFORM
    customer: DÀNH CHO KHÁCH HÀNG A (KHA)
    response_to: "Phản hồi yêu cầu: ..."
    basis: "Căn cứ: ..."
    date: Tháng 10/2026
    version: Phiên bản 1.0. Tài liệu lưu hành nội bộ giữa KHA và Katalyst
    header: Đề xuất giải pháp KHA Data Warehouse trên Hanas Data Platform
Heading: '# 1. TÓM TẮT ĐỀ XUẤT' (sang trang mới), '## 1.1. Bối cảnh', '### 1.1.1. ...' — tự đánh số trong text.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REFERENCE = ROOT / "templates" / "proposal-reference.docx"
FONT = "Roboto Condensed"
HEAD_FILL, BORDER_COLOR = "2E75B6", "BFBFBF"
# (khóa front matter, cỡ chữ, đậm, nghiêng, màu, khoảng cách trước, sau)
COVER = [
    ("title", 20, True, False, "1F3864", 2600, 120),
    ("subtitle", 20, True, False, "1F3864", 0, 400),
    ("customer", 16, True, False, "2E75B6", 0, 1200),
    ("response_to", 11, False, True, None, 0, 120),
    ("basis", 11, False, True, None, 0, 2400),
    ("date", 12, False, False, None, 0, 120),
    ("version", 10, False, False, "7F7F7F", 0, 120),
]


def split_front_matter(text):
    if not text.startswith("---"):
        sys.exit("Thiếu front matter (vendor, title, subtitle, customer, ...) ở đầu file")
    _, fm, body = text.split("---", 2)
    meta = {}
    for line in fm.strip().splitlines():
        key, _, val = line.partition(":")
        meta[key.strip()] = val.strip().strip('"')
    return meta, body


def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def insert_before(anchor, doc, text="", **fmt):
    p = anchor.insert_paragraph_before()
    space_before, space_after, center = fmt.pop("before", 0), fmt.pop("after", 120), fmt.pop("center", True)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before / 20)
    p.paragraph_format.space_after = Pt(space_after / 20)
    if text:
        set_font(p.add_run(text), **fmt)
    return p


def add_cover_and_toc(doc, meta):
    first = doc.paragraphs[0]
    p = insert_before(first, doc, meta.get("vendor", "KATALYST"), size=18, bold=True, color="2E75B6", center=False)
    for key, size, bold, italic, color, before, after in COVER:
        if meta.get(key):
            insert_before(first, doc, meta[key], size=size, bold=bold, italic=italic, color=color,
                          before=before, after=after)
    toc_title = insert_before(first, doc, "MỤC LỤC", size=14, bold=True, after=200)
    toc_title.paragraph_format.page_break_before = True
    toc = insert_before(first, doc)
    run = toc.add_run()
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin", "w:dirty": "true"}),
                      ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "separate"}),
                      ("w:t", None), ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        for k, v in (attr or {}).items():
            el.set(qn(k), v)
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = 'TOC \\h \\o "1-3"'
        if tag == "w:t":
            el.text = "Nhấn F9 (hoặc chuột phải → Update Field) để cập nhật mục lục."
        run._r.append(el)
    return p


def style_table(table):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is not None:  # bảng giãn hết bề rộng trang như mẫu
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", BORDER_COLOR)):
            el.set(qn(k), v)
        borders.append(el)
    margins = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    # schema OOXML bắt buộc thứ tự: ... tblW, tblBorders, ..., tblCellMar, tblLook
    look = tbl_pr.find(qn("w:tblLook"))
    for el in (borders, margins):
        look.addprevious(el) if look is not None else tbl_pr.append(el)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), HEAD_FILL)
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = None
                for r in p.runs:
                    set_font(r, size=10.5, bold=True if i == 0 else None, color="FFFFFF" if i == 0 else None)


def style_body(doc):
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else ""
        if name == "Heading 1":
            p.paragraph_format.page_break_before = True
        if not name.startswith("Heading") and p.alignment is None:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            set_font(r)


def set_header(doc, meta):
    for section in doc.sections:
        for p in section.header.paragraphs:
            p.text = ""
        p = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        set_font(p.add_run(meta.get("vendor", "Katalyst").title()), size=9, bold=True, color="2E75B6")
        set_font(p.add_run("\t" + meta.get("header", "")), size=9, color="7F7F7F")


def review_passed(folder):
    review = Path(folder) / "02-review.md"
    return review.exists() and "KẾT LUẬN: ĐẠT" in review.read_text(encoding="utf-8")


def main():
    args = [a for a in sys.argv[1:] if a != "--draft"]
    draft = "--draft" in sys.argv[1:]
    if len(args) != 2:
        sys.exit(__doc__)
    src, out = Path(args[0]), Path(args[1]).resolve()  # pandoc chạy với cwd=src.parent
    if draft:
        out = out.with_stem(out.stem + "-NHAP")
    elif not review_passed(src.parent):
        sys.exit(f"Chưa xuất: {src.parent / '02-review.md'} chưa có 'KẾT LUẬN: ĐẠT'. "
                 "Chạy vòng review (skill write-proposal bước 7) hoặc thêm --draft cho bản nháp nội bộ.")
    meta, body = split_front_matter(src.read_text(encoding="utf-8"))
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
        f.write(body)
    subprocess.run(["pandoc", f.name, "-f", "gfm", "-o", str(out), f"--reference-doc={REFERENCE}"],
                   check=True, cwd=src.parent)
    Path(f.name).unlink()

    doc = Document(out)
    style_body(doc)
    for t in doc.tables:
        style_table(t)
    add_cover_and_toc(doc, meta)
    set_header(doc, meta)
    doc.save(out)
    print(f"Đã ghi {out} ({len(doc.paragraphs)} đoạn, {len(doc.tables)} bảng). Mở Word, nhấn F9 để cập nhật mục lục.")


if __name__ == "__main__":
    main()
